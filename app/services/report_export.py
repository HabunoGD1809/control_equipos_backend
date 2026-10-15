"""Shared branded exports. User supplied values are always exported as data."""
import csv
import os
import math
from datetime import date, datetime, timezone
from pathlib import Path
from decimal import Decimal
from fpdf import FPDF
from fpdf.fonts import FontFace
from fpdf.enums import MethodReturnValue
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

NAVY = "17324D"
TEAL = "007F86"


def display(value):
    if value is None:
        return ""
    if hasattr(value, "value"):
        return str(value.value)
    if isinstance(value, datetime):
        return (value.astimezone(timezone.utc) if value.tzinfo else value).strftime("%d/%m/%Y %H:%M")
    return str(value)


def spreadsheet_value(value):
    if isinstance(value, datetime) and value.tzinfo:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    if isinstance(value, Decimal):
        return float(value)
    return value if value is not None else ""


def generar_csv(path, rows):
    with Path(path).open("w", newline="", encoding="utf-8-sig") as stream:
        if not rows:
            csv.writer(stream).writerow(["No hay datos para el período seleccionado."])
            return
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        for row in rows:
            writer.writerow({key: value.astimezone(timezone.utc).isoformat() if isinstance(value, datetime) and value.tzinfo else spreadsheet_value(value) for key, value in row.items()})


def fonts():
    candidates = [
        (os.getenv("REPORT_FONT_PATH"), os.getenv("REPORT_FONT_BOLD_PATH")),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ]
    for regular, bold in candidates:
        if regular and bold and Path(regular).is_file() and Path(bold).is_file():
            return regular, bold
    raise RuntimeError("Instale las fuentes DejaVu o configure REPORT_FONT_PATH y REPORT_FONT_BOLD_PATH.")


class CompanyPDF(FPDF):
    def __init__(self, company, title, period, orientation="L"):
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.company, self.title, self.period = company or {}, title, period
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(True, margin=16)
        regular, bold = fonts()
        self.add_font("Report", "", regular)
        self.add_font("Report", "B", bold)
        self.alias_nb_pages()

    def header(self):
        logo = self.company.get("logo")
        if logo and Path(logo).is_file():
            self.image(str(logo), x=12, y=12, w=25, h=19, keep_aspect_ratio=True)
        self.set_xy(43 if logo else 12, 12)
        self.set_font("Report", "B", 14)
        self.set_text_color(23, 50, 77)
        self.multi_cell(0, 7, self.company.get("nombre", "ControlEquipos"), new_x="LMARGIN", new_y="NEXT")
        self.set_x(43 if logo else 12)
        self.set_font("Report", "", 8)
        contact = " · ".join(str(self.company[key]) for key in ["identificacion", "telefono", "email"] if self.company.get(key))
        self.multi_cell(0, 5, contact, new_x="LMARGIN", new_y="NEXT")
        self.set_y(max(35, self.get_y() + 3))
        self.set_font("Report", "B", 13)
        self.multi_cell(0, 7, self.title, new_x="LMARGIN", new_y="NEXT")
        self.set_font("Report", "", 8)
        self.set_text_color(90, 105, 120)
        self.multi_cell(0, 5, self.period, new_x="LMARGIN", new_y="NEXT")
        self.ln(4)
        self.content_start_y = self.get_y()

    def footer(self):
        self.set_y(-12)
        self.set_draw_color(205, 215, 222)
        self.line(12, self.get_y() - 2, self.w - 12, self.get_y() - 2)
        self.set_font("Report", "", 7)
        self.set_text_color(90, 105, 120)
        self.cell(0, 5, f"Fechas y horas en UTC  |  Generado {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC  |  Página {self.page_no()}/{{nb}}", align="R")


def generar_pdf(path, rows, title, company=None, period=""):
    pdf = CompanyPDF(company, title, period)
    pdf.add_page()
    pdf.set_font("Report", "", 8)
    pdf.set_text_color(23, 50, 77)
    pdf.cell(0, 6, f"Total de registros: {len(rows)}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    if not rows:
        pdf.cell(0, 10, "No hay datos para el período seleccionado.")
    else:
        columns = list(rows[0])
        weights = [min(30, max(10, len(key), max(len(display(row.get(key))) for row in rows[:200]) ** .65 * 3)) for key in columns]
        pdf.set_font("Report", "", 7 if len(columns) > 8 else 8)
        # Reserve room for the longest header word before distributing space to data.
        # Narrow numeric columns must still fit labels such as Entrada and Saldo.
        pdf.set_font("Report", "B", pdf.font_size_pt)
        minimums = [max(9, max((pdf.get_string_width(word) for word in key.split()), default=0) + 3.5) for key in columns]
        for index, key in enumerate(columns):
            labels = [str(row[key]) for row in rows[:200] if isinstance(row.get(key), str) and 0 < len(row[key]) <= 14 and not any(character.isspace() for character in row[key])]
            if labels:
                minimums[index] = max(minimums[index], max(pdf.get_string_width(value) for value in labels) + 3.5)
        # Dates must fit as a whole; splitting the year makes a dense kardex hard to read.
        for index, key in enumerate(columns):
            dates = [display(row[key]).split()[0] for row in rows[:200] if isinstance(row.get(key), date)]
            if dates:
                minimums[index] = max(minimums[index], max(pdf.get_string_width(value) for value in dates) + 3.5)
        remaining = max(0, pdf.epw - sum(minimums))
        weights = [minimum + remaining * weight / sum(weights) for minimum, weight in zip(minimums, weights)]
        pdf.set_font("Report", "", pdf.font_size_pt)
        write_pdf_table(pdf, columns, rows, weights)
    pdf.output(str(path))


def write_pdf_table(pdf, columns, rows, weights, line_height=4.5, padding=1.5):
    """Wrap long values without losing lines when a table spans several pages."""
    widths = [pdf.epw * weight / sum(weights) - 2 * padding for weight in weights]
    auto_break, bottom_margin = pdf.auto_page_break, pdf.b_margin
    font_size = pdf.font_size_pt
    pdf.set_auto_page_break(False, margin=bottom_margin)
    try:
        pdf.set_font("Report", "B", font_size)
        headings = [pdf.multi_cell(width, line_height, key, dry_run=True, output=MethodReturnValue.LINES) or [""] for key, width in zip(columns, widths)]
        header_height = max(map(len, headings)) * line_height + 2 * padding
        max_lines = max(1, min(20, int((pdf.page_break_trigger - pdf.content_start_y - header_height - 2 * padding - 1) / line_height)))
        pdf.set_font("Report", "", font_size)
        values_to_export = []
        for record in rows:
            wrapped = [pdf.multi_cell(width, line_height, display(record.get(key)), dry_run=True, output=MethodReturnValue.LINES) or [""] for key, width in zip(columns, widths)]
            for offset in range(0, max(map(len, wrapped)), max_lines):
                lines = [value[offset:offset + max_lines] for value in wrapped]
                values_to_export.append((["\n".join(value) for value in lines], max(map(len, lines)) * line_height + 2 * padding))
    finally:
        pdf.set_auto_page_break(auto_break, margin=bottom_margin)

    def render(batch):
        # Explicit pagination keeps the document header outside table measurement.
        pdf.set_auto_page_break(False, margin=bottom_margin)
        try:
            with pdf.table(col_widths=weights, line_height=line_height, text_align="LEFT", padding=padding,
                           headings_style=FontFace(emphasis="BOLD", color=(255,255,255), fill_color=(23,50,77)),
                           cell_fill_color=(242,246,249), cell_fill_mode="ROWS", borders_layout="MINIMAL") as table:
                for values in [columns, *batch]:
                    table_row = table.row()
                    for value in values:
                        table_row.cell(value)
        finally:
            pdf.set_auto_page_break(auto_break, margin=bottom_margin)

    batch, height = [], header_height
    for values, row_height in values_to_export:
        if pdf.get_y() + height + row_height > pdf.page_break_trigger - .5:
            if batch:
                render(batch)
            pdf.add_page()
            batch, height = [], header_height
        batch.append(values)
        height += row_height
    if batch:
        render(batch)


def company_identity(db):
    from app.models.empresa import Empresa
    from app.core.storage import UPLOAD_DIR
    empresa = db.get(Empresa, 1)
    company = {key: getattr(empresa, key) for key in ["nombre", "identificacion", "direccion", "telefono", "email"]} if empresa else {"nombre": "ControlEquipos"}
    if empresa and empresa.logo_path:
        logo_dir = (UPLOAD_DIR / "branding").resolve()
        logo = (logo_dir / empresa.logo_path).resolve()
        if logo.parent == logo_dir and logo.is_file():
            company["logo"] = str(logo)
    return company


def generar_excel(path, rows, title="Reporte de inventario", company=None, period=""):
    company = company or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte"
    sheet.sheet_view.showGridLines = False
    columns = list(rows[0]) if rows else ["Aviso"]
    last = max(len(columns), 4)
    has_logo = company.get("logo") and Path(company["logo"]).is_file()
    for row_number, text in [(1,company.get("nombre", "ControlEquipos")), (2,title), (3,period), (4,f"Registros: {len(rows)}. Fechas y horas en UTC. Generado {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC")]:
        sheet.merge_cells(start_row=row_number,start_column=1,end_row=row_number,end_column=last-1 if row_number==1 and has_logo else last)
        cell = sheet.cell(row_number,1,spreadsheet_value(text))
        cell.font=Font(name="Calibri",size=18 if row_number==1 else 11,bold=row_number<3,color="FFFFFF" if row_number==1 else NAVY)
        cell.fill=PatternFill("solid",fgColor=NAVY if row_number==1 else "F2F6F9")
        cell.alignment=Alignment(vertical="center",wrap_text=True)
        sheet.row_dimensions[row_number].height=30 if row_number==1 else 24
    if has_logo:
        logo=ExcelImage(company["logo"])
        ratio=min(90/logo.width,32/logo.height)
        logo.width,logo.height=logo.width*ratio,logo.height*ratio
        sheet.add_image(logo,f"{get_column_letter(last)}1")
        sheet.column_dimensions[get_column_letter(last)].width=18
    for col,key in enumerate(columns,1):
        cell=sheet.cell(6,col,key)
        cell.font=Font(name="Calibri",bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor=TEAL)
        cell.alignment=Alignment(wrap_text=True,vertical="center")
    for row_number,data in enumerate(rows or [{"Aviso":"No hay datos para el período seleccionado."}],7):
        for col,key in enumerate(columns,1):
            value=spreadsheet_value(data.get(key))
            cell=sheet.cell(row_number,col,value)
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            cell.font=Font(name="Calibri",size=11,color=NAVY)
            if isinstance(value,datetime):
                cell.number_format='dd/mm/yyyy hh:mm'
            if isinstance(value,(int,float,Decimal)):
                cell.number_format='#,##0.00' if isinstance(value,(float,Decimal)) else '#,##0'
            if row_number%2:
                cell.fill=PatternFill("solid",fgColor="F2F6F9")
        sheet.row_dimensions[row_number].height=32
    for col,key in enumerate(columns,1):
        lengths=[len(display(row.get(key))) for row in rows[:200]]
        sheet.column_dimensions[get_column_letter(col)].width=min(48,max(15,len(key)+3,max(lengths,default=0)+2))
    for row_number, data in enumerate(rows or [{"Aviso": "No hay datos para el período seleccionado."}], 7):
        lines = max(sum(max(1, math.ceil(len(line) / max(10, sheet.column_dimensions[get_column_letter(col)].width - 2))) for line in display(data.get(key)).split("\n")) for col, key in enumerate(columns, 1))
        # Excel's maximum row height is 409 points; full text remains available in each cell.
        sheet.row_dimensions[row_number].height = min(409, max(32, lines * 15 + 6))
    if rows:
        table=Table(displayName="DatosReporte",ref=f"A6:{get_column_letter(len(columns))}{len(rows)+6}")
        table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes="A7"
    sheet.auto_filter.ref=f"A6:{get_column_letter(len(columns))}{sheet.max_row}"
    sheet.print_title_rows="1:6"
    sheet.print_options.horizontalCentered=True
    sheet.page_setup.orientation="landscape"
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1
    sheet.page_setup.fitToHeight=0
    sheet.print_area=f"A1:{get_column_letter(last)}{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    workbook.save(path)
