import csv
import io
import re
from uuid import uuid4
from decimal import Decimal
from datetime import datetime, timezone
import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from app.services.report_export import generar_csv, generar_excel, generar_pdf
from app.core.password import get_password_hash, verify_password
from app.core.security import create_access_token, create_refresh_token, decode_access_token, decode_refresh_token
from app.core.storage import save_upload_file


def test_unicode_passwords_are_not_truncated_at_bcrypt_limit():
    password = "ClaveÁ" * 20
    encoded = get_password_hash(password)
    assert verify_password(password, encoded)
    assert not verify_password(password + "diferente", encoded)


def test_refresh_tokens_cannot_authenticate_as_access_tokens():
    user_id = uuid4()
    access, refresh = create_access_token(user_id), create_refresh_token(user_id)
    assert decode_access_token(access).sub == str(user_id)
    assert decode_refresh_token(refresh).sub == str(user_id)
    assert decode_access_token(refresh) is None
    assert decode_refresh_token(access) is None
    assert create_refresh_token(user_id) != refresh


def test_exports_preserve_data_and_prevent_formula_execution(tmp_path):
    rows = [{"Equipo": "Álvarez", "Nota": "=HYPERLINK(\"https://invalid.example\")", "Costo": Decimal("123.45"), "Fecha": datetime.now(timezone.utc)}]
    generar_csv(tmp_path/'datos.csv', rows)
    with (tmp_path/'datos.csv').open(encoding='utf-8-sig', newline='') as stream:
        exported = list(csv.DictReader(stream))
    assert exported[0]['Equipo'] == 'Álvarez'
    assert exported[0]['Nota'].startswith("'=")
    generar_excel(tmp_path/'datos.xlsx', rows, 'Equipos', {'nombre': 'Informática Álvarez'})
    sheet = load_workbook(tmp_path/'datos.xlsx').active
    assert sheet['B7'].data_type == 's'
    assert sheet['B7'].value.startswith("'=")
    assert sheet['C7'].value == 123.45
    assert sheet['D7'].value.tzinfo is None
    assert sheet.freeze_panes == 'A7'
    assert len(sheet.tables) == 1


def test_pdf_paginates_long_cells_without_truncation_or_failure(tmp_path):
    path = tmp_path/'reporte.pdf'
    rows = [{"Equipo": "Servidor", "Observación": "Observación técnica extensa con acentos. " * 500}]
    generar_pdf(path, rows, 'Auditoría', {'nombre': 'Informática Álvarez'}, 'Período de prueba')
    content = path.read_bytes()
    assert content.startswith(b'%PDF-')
    assert content.rstrip().endswith(b'%%EOF')
    assert len(re.findall(rb'/Type /Page\b', content)) >= 2


@pytest.mark.asyncio
async def test_disguised_executable_is_rejected_as_image(tmp_path):
    file = UploadFile(filename='imagen.png', file=io.BytesIO(b'MZ' + b'fake image'))
    with pytest.raises(HTTPException) as error:
        await save_upload_file(file)
    assert error.value.status_code == 415
